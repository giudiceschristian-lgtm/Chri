"""Controlli sui .docx generati: si aprono, il corpo è XML valido, intestazione/piè di pagina/filigrana/stili
identici al modello, conversione in PDF con LibreOffice, coerenza FAQ ↔ piano."""
import json, zipfile, subprocess, os, re, sys, hashlib
from lxml import etree
import docx
HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, '..', 'sito', 'FSM_ALIMENTAZIONE.docx')
rep = json.load(open(os.path.join(HERE, 'out', 'report.json')))
t = zipfile.ZipFile(TPL)
KEEP = ['word/header1.xml', 'word/footer1.xml', 'word/_rels/header1.xml.rels', 'word/media/image1.png', 'word/media/image2.png',
        'word/styles.xml', 'word/theme/theme1.xml', 'word/settings.xml', 'word/numbering.xml', 'word/fontTable.xml', '[Content_Types].xml']
h = lambda b: hashlib.md5(b).hexdigest()
results = []
for r in rep:
    if r.get('error'): results.append((r['name'], ['ERRORE generazione']))
    f = r['file']; issues = []; info = {}
    z = zipfile.ZipFile(f)
    for k in KEEP:
        if h(z.read(k)) != h(t.read(k)): issues.append('parte modificata: ' + k)
    body = z.read('word/document.xml')
    etree.fromstring(body)  # XML valido
    sect_tpl = re.search(rb'<w:sectPr.*?</w:sectPr>', t.read('word/document.xml'), re.S).group(0)
    if sect_tpl not in body: issues.append('sectPr (rimandi a intestazione/piè di pagina) non conservato')
    d = docx.Document(f)
    info['paragrafi'] = len(d.paragraphs); info['tabelle'] = len(d.tables)
    info['intestazione_immagini'] = len(d.sections[0].header._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'))
    info['piede'] = ' '.join(p.text for p in d.sections[0].footer.paragraphs).strip()
    if info['intestazione_immagini'] < 2: issues.append('logo o filigrana mancanti')
    out = subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', os.path.join(HERE, 'out'), f], capture_output=True, timeout=180)
    pdf = f[:-5] + '.pdf'
    if not os.path.exists(pdf): issues.append('conversione PDF fallita')
    else:
        pg = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True).stdout
        info['pagine'] = int(re.search(r'Pages:\s+(\d+)', pg).group(1))
    # coerenza FAQ / piano
    text = r['docText']; faq = ' '.join(r['faq'])
    water = re.search(r'ACQUA: MIN ([\d,]+)LT', text)
    if water and water.group(1) not in faq: issues.append('FAQ acqua non coerente')
    has_leg = any(i in r['items'] for i in ['lenticchie_secche','ceci_secchi','fagioli_cannellini_secchi','fagioli_borlotti_secchi','lenticchie_rosse_decorticate','ceci_in_barattolo_sgocciolati','fagioli_in_barattolo_sgocciolati'])
    if has_leg != ('legumi?' in faq): issues.append('FAQ legumi incoerente (legumi nel piano: %s)' % has_leg)
    multi = r['type'] == 'multi'
    if ('Come combino le opzioni?' in faq) != multi: issues.append('FAQ "come combinare" incoerente col tipo di piano')
    if ('ESEMPI PASTI' in text) and not multi: issues.append('Esempi pasti nel piano giorno per giorno')
    if multi and 'ESEMPI PASTI' not in text and 'PASSO' in text: issues.append('Esempi pasti mancanti')
    if 'NESSUN PASTO LIBERO' in text and 'pasto libero' in faq.lower() and 'non sono previsti pasti liberi' not in faq: issues.append('FAQ pasto libero incoerente')
    m = re.search(r'PASTO LIBERO: (\w+) A (\w+)', text)
    if m and (m.group(1).lower() not in faq.lower() or m.group(2).lower() not in faq.lower()): issues.append('FAQ pasto libero: giorno/pasto diverso')
    if 'lattosio' in r['name'].lower() and 'intollerante al lattosio' not in faq: issues.append('FAQ lattosio mancante')
    if 'intollerante al lattosio' in faq and 'lattosio' not in r['name'].lower(): issues.append('FAQ lattosio senza flag')
    if len(r['faq']) > 12: issues.append('più di 12 FAQ')
    if 'Si declinano ogni responsabilità' not in text: issues.append('disclaimer mancante')
    results.append((r['name'], issues, info))
for name, issues, *rest in results:
    print('###', name); print('   ', rest[0] if rest else ''); print('    esito:', 'OK' if not issues else issues)
